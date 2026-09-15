import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook

from tools import _run_three_sheets as runner
from tools.exception_queue import build_exception_queue
from tools.execution_journal import ExecutionJournal, JournalError
from tools.module_planner import build_module_plan
from tools.execution_gate import validate_execution_contract


def _make_cases(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "模块A"
    sheet.append(["1级目录", "用例名称", "优先级", "入口", "步骤名称", "前置条件", "操作描述", "参数", "预期结果"])
    sheet.append(["行情", "TC_打开模块", "P0", "首页", "进入", "竖屏", "点击模块", "", "显示模块首页"])
    sheet.append(["行情", "TC_查看列表", "P1", "模块A", "查看", "", "查看列表", "", "显示列表"])
    sheet2 = workbook.create_sheet("模块B")
    sheet2.append(["用例名称", "操作描述", "预期结果"])
    sheet2.append(["TC_刷新", "点击刷新", "数据更新"])
    workbook.save(path)


def test_module_plan_reads_once_and_keeps_row_order(tmp_path):
    source = tmp_path / "cases.xlsx"
    _make_cases(source)

    plan = build_module_plan(source, ["模块A", "模块B"])

    assert plan["planning_scope"] == "module_page_group"
    assert plan["execution_scope"] == "single_excel_row"
    assert plan["page_batching_allowed"] is True
    assert plan["expected_count"] == 3
    assert [case["row"] for case in plan["modules"][0]["cases"]] == [2, 3]
    assert plan["modules"][0]["cases"][0]["action"] == "点击模块"
    assert plan["execution_manifest"]["selected_cases"][2]["source_order"] == 3


def test_page_groups_include_hierarchy_and_do_not_merge_same_page_name(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "行情"
    sheet.append(["一级目录", "二级目录", "三级目录", "用例名称", "入口", "操作描述", "预期结果"])
    sheet.append(["行情", "沪深京", "沪深A股", "沪深详情-1", "沪深A股", "点击第一条", "显示详情"])
    sheet.append([None, None, None, "沪深详情-2", "沪深A股", "横屏", "横屏详情"])
    sheet.append(["行情", "港股", "港股", "港股详情-1", "港股", "点击第一条", "显示详情"])
    source = tmp_path / "context.xlsx"
    workbook.save(source)

    plan = build_module_plan(source, ["行情"])
    cases = plan["modules"][0]["cases"]

    assert cases[0]["navigation_context"]["level_3"] == "沪深A股"
    assert cases[1]["navigation_context"]["level_3"] == "沪深A股"
    assert cases[0]["page_group_id"] == cases[1]["page_group_id"]
    assert cases[1]["page_group_id"] != cases[2]["page_group_id"]
    assert cases[2]["navigation_context"]["level_2"] == "港股"
    assert plan["modules"][0]["page_group_count"] == 2


def test_summarize_records_page_observation_without_action_echo():
    actual = runner.summarize([{"id": "query_page"}, {"text": "结果列表"}])

    assert actual == "当前页面观察：query_page、结果列表"
    assert "点击查询" not in actual


def test_build_actual_uses_trace_steps_and_screenshot_visible_text():
    actual = runner.build_actual(
        [
            {"type": "tap", "target": "顶部港股", "result": "success"},
            {"type": "evidence", "target": "shot.png", "result": "success"},
        ],
        [{"text": "港股"}, {"text": "AH股"}, {"id": "action_bar_root"}],
        setup_ok=True,
        action_ok=True,
        judgment_status="✅通过",
        judgment_reason="截图中的港股标题和AH股区域支持该结论",
    )

    assert "AI执行步骤：" in actual
    assert "点击已识别控件“顶部港股”" in actual
    assert "截图可见文字/标题：港股、AH股" in actual
    assert "判断理由：判定为✅通过。截图中的港股标题和AH股区域支持该结论" in actual
    assert "action_bar_root" not in actual
    assert "行情界面，点击“港股”tab页" not in actual


def test_row_target_contract_rejects_detail_and_accepts_a_share_list(monkeypatch):
    case = {
        "case_name": "TC_表头字段排序",
        "step_name": "横屏模式",
        "precondition": "横屏模式",
        "action": "进入其他--个股-上证A股横屏\n点击表头字段",
    }
    monkeypatch.setattr(runner, "orientation_matches", lambda landscape: landscape)

    contract = runner.target_page_contract("其他", 33, case)
    detail_page = [
        {"id": "page_queue_nav_bar"},
        {"id": "navi_animation_label"},
        {"id": "backButton"},
        {"text": "深纺织B"},
    ]
    a_share_list = [
        {"id": "hx_page_title_bar"},
        {"id": "table"},
        {"text": "返回"},
        {"text": "上证A股"},
    ]

    assert not contract.predicate(detail_page)
    assert contract.predicate(a_share_list)


def test_hk_home_uses_dedicated_surface_and_allows_scrolled_home():
    home = [
        {"id": "title_bar_middle"},
        {"id": "ganggu_page"},
        {"id": "titlebar_left_layout"},
        {"text": "港股"},
        {"text": "沪深京"},
        {"text": "其他"},
        {"text": "恒生指数"},
        {"text": "国企指数"},
        {"text": "沪、深港通"},
    ]
    scrolled_home = [item for item in home if item.get("text") not in {"恒生指数", "国企指数", "沪、深港通"}]
    incomplete = [item for item in home if item.get("text") != "港股"]

    assert runner.is_market_shell(home)
    assert runner.is_market_home(home, "港股")
    assert runner.is_market_home(scrolled_home, "港股")
    assert not runner.is_market_home(incomplete, "港股")


def test_enter_market_home_prefers_top_tab_over_same_text_title(monkeypatch):
    calls = []
    shell = [
        {"id": "title_bar_middle"},
        {"text": "沪深京"},
        {"text": "港股"},
        {"text": "其他"},
    ]
    monkeypatch.setattr(runner, "screen_elements", lambda: shell)
    monkeypatch.setattr(runner, "tap_xy", lambda events, x, y, target: calls.append((x, y)) or True)
    monkeypatch.setattr(runner, "tap_text", lambda *args: calls.append("text") or False)
    monkeypatch.setattr(runner, "wait_for_page", lambda predicate, description: True)

    assert runner.enter_market_home([], "港股")
    assert calls == [(756, 277)]


def test_target_page_guard_rechecks_after_navigation(monkeypatch):
    calls = []
    wrong_page = [{"id": "page_queue_nav_bar"}]
    target_page = [{"id": "table"}, {"text": "返回"}, {"text": "上证A股"}]
    monkeypatch.setattr(runner, "screen_elements", lambda: wrong_page)

    def navigate(events):
        calls.append("navigate")
        return True

    def wait_for_target(predicate, description):
        calls.append("recheck")
        return predicate(target_page)

    monkeypatch.setattr(runner, "wait_for_page", wait_for_target)
    contract = runner.PageContract(
        "上证A股列表页",
        lambda elements: runner.is_named_list_page(elements, ("上证A股",)),
        navigate,
    )

    events = []
    assert runner.ensure_target_page(events, contract)
    assert calls == ["navigate", "recheck"]
    assert [item["result"] for item in events if item["type"] == "assert"] == ["failed", "success"]


def test_retest_queue_resolves_only_requested_rows(tmp_path):
    queue_path = tmp_path / "retest_queue.json"
    queue_path.write_text(
        json.dumps(
            {
                "cases": [
                    {
                        "retest_id": "模块A!3",
                        "retest_order": 1,
                        "sheet": "模块A",
                        "row": 3,
                        "case_id": "模块A-row-003",
                        "initial_status": "🟡待验证",
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    entries = runner._load_retest_queue(queue_path)
    plan = {
        "modules": [
            {
                "sheet": "模块A",
                "page_groups": [
                    {
                        "page_group_id": "模块A-page-group-001",
                        "page_group_key": "模块A|行情",
                        "cases": [
                            {
                                "sheet": "模块A",
                                "row": 2,
                                "case_id": "模块A-row-002",
                                "source_order": 1,
                            },
                            {
                                "sheet": "模块A",
                                "row": 3,
                                "case_id": "模块A-row-003",
                                "source_order": 2,
                            },
                        ],
                    }
                ],
            }
        ]
    }

    items = runner._retest_execution_items(plan, entries)

    assert [item[2]["row"] for item in items] == [3]
    assert items[0][1]["page_group_id"] == "模块A-page-group-001"


def test_retest_queue_carries_first_pass_agent_binding(tmp_path):
    queue_path = tmp_path / "bound_retest_queue.json"
    queue_path.write_text(
        json.dumps(
            {
                "agent_binding": {
                    "planner": {"agent": "Trae", "model": "model-a"},
                    "retester": {"agent": "Trae", "model": "model-a"},
                },
                "cases": [
                    {
                        "sheet": "模块A",
                        "row": 3,
                        "case_id": "模块A-row-003",
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    entries = runner._load_retest_queue(queue_path)

    assert entries[0]["agent_binding"]["planner"]["agent"] == "Trae"


def test_retest_queue_rejects_duplicate_rows(tmp_path):
    queue_path = tmp_path / "duplicate_queue.json"
    queue_path.write_text(
        json.dumps(
            {
                "cases": [
                    {"sheet": "模块A", "row": 3, "case_id": "模块A-row-003"},
                    {"sheet": "模块A", "row": 3, "case_id": "模块A-row-003"},
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    try:
        runner._load_retest_queue(queue_path)
    except ValueError as exc:
        assert "重复用例" in str(exc)
    else:
        raise AssertionError("duplicate retest row should be rejected")


def test_blocked_retest_command_disables_recursive_retest(tmp_path):
    command = runner._blocked_retest_command(
        app_slug="zhongyuan",
        source=tmp_path / "cases.xlsx",
        profile=tmp_path / "profile.yaml",
        device="device-1",
        output=tmp_path / "blocked-retest",
        queue_path=tmp_path / "blocked_retest_queue.json",
        action_plan=str(tmp_path / "agent_action_plan.json"),
        legacy_deterministic=False,
        resume=False,
    )

    assert "--retest-queue" in command
    assert "--no-auto-retest-blocked" in command
    assert "--resume" not in command


def test_blocked_retest_command_can_select_stepwise_llm_mode(tmp_path):
    command = runner._blocked_retest_command(
        app_slug="zhongyuan",
        source=tmp_path / "cases.xlsx",
        profile=tmp_path / "profile.yaml",
        device="device-1",
        output=tmp_path / "blocked-retest",
        queue_path=tmp_path / "blocked_retest_queue.json",
        action_plan=str(tmp_path / "agent_action_plan.json"),
        legacy_deterministic=False,
        llm_retest=True,
        resume=False,
    )

    assert "--llm-retest" in command
    assert "--no-auto-retest-blocked" in command
    assert "--action-plan" not in command


def test_blocked_retest_runs_once_and_merges(tmp_path, monkeypatch):
    def case(row, status, result, evidence):
        reason = f"判定为{status}。{result}"
        return {
            "module": "股指",
            "sheet": "股指",
            "row": row,
            "case_id": f"股指-row-{row:03d}",
            "case_name": f"TC-{row}",
            "source_order": row - 1,
            "execution_order": row - 1,
            "status": status,
            "actual": (
                "AI执行步骤：\n1. 执行本行操作\n"
                f"操作结果：\n{result}\n"
                f"判断理由：{reason}"
            ),
            "judgment_reason": reason,
            "action_trace": [
                {
                    "type": "tap",
                    "target": f"row-{row}",
                    "result": "failed" if status == "⛔阻塞" else "success",
                }
            ],
            "evidence": [evidence],
        }

    initial = {
        "schema_version": "2.0",
        "cases": [
            case(2, "⛔阻塞", "首轮未找到控件", "first/row-2.png"),
            case(3, "✅通过", "首轮页面正确", "first/row-3.png"),
        ],
    }
    def fake_run(command, *, cwd, check):
        assert check is False
        assert cwd == str(runner.PROJECT_ROOT)
        child_output = Path(command[command.index("--output") + 1])
        child_output.mkdir(parents=True, exist_ok=True)
        retest = {
            "schema_version": "2.0",
            "cases": [
                case(2, "✅通过", "第二轮重新进入后控件可见", "second/row-2.png")
            ],
        }
        (child_output / "retest_execution.json").write_text(
            json.dumps(retest, ensure_ascii=False),
            encoding="utf-8",
        )
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(runner.subprocess, "run", fake_run)
    merged, merged_path, summary, exit_code = runner._run_blocked_retest_once(
        initial,
        output=tmp_path,
        app_slug="zhongyuan",
        source=tmp_path / "cases.xlsx",
        profile=tmp_path / "profile.yaml",
        device="device-1",
        action_plan=str(tmp_path / "agent_action_plan.json"),
        legacy_deterministic=False,
    )

    assert exit_code == 0
    assert summary["planned"] == 1
    assert summary["completed"] == 1
    assert merged_path == tmp_path / "execution_records.retested.json"
    assert merged["cases"][0]["status"] == "✅通过"
    assert [attempt["phase"] for attempt in merged["cases"][0]["attempts"]] == [
        "batch",
        "single_case_retest",
    ]
    assert merged["cases"][1]["status"] == "✅通过"


def test_detail_case_name_is_used_only_when_action_needs_detail():
    click_stock = {
        "case_name": "TC_列表跳转个股详情界面",
        "action": "进入看资金--概念\n点击列表中任意个股",
        "precondition": "横屏模式",
    }
    close_detail = {
        "case_name": "TC_个股详情界面返回概念板块列表",
        "action": "进入看资金--概念\n点击个股分时图右上角的X按钮",
        "precondition": "横屏模式",
    }

    assert not runner._target_is_detail(click_stock)
    assert runner._target_is_detail(close_detail)


def test_journal_persists_each_row_and_resumes(tmp_path):
    manifest = {
        "mode": "full",
        "expected_count": 1,
        "selected_cases": [{"sheet": "模块A", "row": 2, "case_id": "模块A-row-002"}],
    }
    journal = ExecutionJournal(tmp_path, manifest)
    journal.append(
        {
            "sheet": "模块A",
            "row": 2,
            "case_id": "模块A-row-002",
            "source_order": 1,
            "execution_order": 1,
            "status": "⛔阻塞",
            "blocked_reason": "控件未找到",
            "actual": "页面未出现目标控件",
            "action_trace": [{"type": "tap", "target": "目标", "result": "not_found"}],
            "evidence": ["shots/模块A-row-002.png"],
        }
    )

    state = json.loads((tmp_path / "execution_state.json").read_text(encoding="utf-8"))
    assert state["completed_count"] == 1
    assert state["last_completed"]["row"] == 2
    assert (tmp_path / "execution_records.jsonl").exists()
    assert journal.finalize().exists()
    assert ExecutionJournal(tmp_path, manifest, resume=True).has_case("模块A", 2)


def test_journal_rejects_duplicate_case(tmp_path):
    manifest = {
        "mode": "full",
        "expected_count": 1,
        "selected_cases": [{"sheet": "模块A", "row": 2, "case_id": "模块A-row-002"}],
    }
    journal = ExecutionJournal(tmp_path, manifest)
    record = {
        "sheet": "模块A",
        "row": 2,
        "case_id": "模块A-row-002",
        "source_order": 1,
        "execution_order": 1,
        "status": "⛔阻塞",
    }
    journal.append(record)
    try:
        journal.append(record)
    except JournalError as exc:
        assert "重复写入用例" in str(exc)
    else:
        raise AssertionError("duplicate case should be rejected")


def test_exception_queue_is_compact_and_excludes_pass_cases():
    queue = build_exception_queue(
        {
            "cases": [
                {"sheet": "模块A", "row": 2, "case_id": "A-2", "status": "⛔阻塞", "actual": "未找到", "evidence": ["a.png"]},
                {"sheet": "模块A", "row": 3, "case_id": "A-3", "status": "✅通过", "actual": "页面显示列表"},
            ]
        }
    )

    assert queue["review_scope"] == "module"
    assert queue["exception_count"] == 1
    assert queue["cases"][0]["exception_id"] == "A-2"


def test_row_scoped_gate_rejects_observe_only_pass_and_missing_order():
    manifest = {
        "mode": "full",
        "execution_scope": "single_excel_row",
        "expected_count": 1,
        "selected_cases": [{"sheet": "模块A", "row": 2, "case_id": "模块A-row-002"}],
    }
    record = {
        "sheet": "模块A",
        "row": 2,
        "case_id": "模块A-row-002",
        "status": "✅通过",
        "action_mode": "observe",
        "actual": "页面显示列表",
        "evidence": ["shots/a.png"],
        "action_trace": [{"type": "observe", "result": "success"}],
    }

    errors = validate_execution_contract([record], manifest)

    assert any("缺少 source_order" in error for error in errors)
    assert any("缺少 execution_order" in error for error in errors)
    assert any("observe-only" in error for error in errors)


def test_module_session_cold_starts_once_then_soft_resets(monkeypatch):
    calls = {"rotate": 0, "launch": 0, "soft_reset": 0}

    def fake_rotate(events, landscape):
        calls["rotate"] += 1
        return True

    def fake_launch(events):
        calls["launch"] += 1
        return True

    def fake_soft_reset(events, sheet_name):
        calls["soft_reset"] += 1
        return True

    monkeypatch.setattr(runner, "rotate", fake_rotate)
    monkeypatch.setattr(runner, "launch_market", fake_launch)
    monkeypatch.setattr(runner, "soft_reset_to_module_root", fake_soft_reset)

    session = runner.ModuleSession()
    assert runner.ensure_module_state([], "港股", session)
    assert runner.ensure_module_state([], "港股", session)

    assert calls == {"rotate": 2, "launch": 1, "soft_reset": 1}
    assert session.cold_start_count == 1
    assert session.soft_reset_count == 1
    assert session.recovery_restart_count == 0


def test_page_group_reuses_navigation_without_module_reset(monkeypatch):
    calls = {"rotate": 0, "module": 0}

    monkeypatch.setattr(runner, "rotate", lambda events, landscape: calls.__setitem__("rotate", calls["rotate"] + 1) or True)
    def fake_module_state(events, sheet, session):
        calls["module"] += 1
        session.active_sheet = sheet
        return True

    monkeypatch.setattr(runner, "ensure_module_state", fake_module_state)

    session = runner.ModuleSession()
    first = runner.ensure_page_group_state([], "港股", "g1", "港股|行情|港股", session)
    second = runner.ensure_page_group_state([], "港股", "g1", "港股|行情|港股", session)

    assert first and second
    assert calls["module"] == 1
    assert calls["rotate"] == 1
    assert session.page_group_reuse_count == 1


def test_module_session_uses_one_bounded_cold_start_when_soft_reset_fails(monkeypatch):
    calls = {"launch": 0, "soft_reset": 0}

    monkeypatch.setattr(runner, "rotate", lambda events, landscape: True)

    def fake_launch(events):
        calls["launch"] += 1
        return True

    def fake_soft_reset(events, sheet_name):
        calls["soft_reset"] += 1
        return False

    monkeypatch.setattr(runner, "launch_market", fake_launch)
    monkeypatch.setattr(runner, "soft_reset_to_module_root", fake_soft_reset)

    session = runner.ModuleSession(active_sheet="港股")
    assert runner.ensure_module_state([], "港股", session)

    assert calls == {"launch": 1, "soft_reset": 1}
    assert session.cold_start_count == 1
    assert session.soft_reset_count == 0
    assert session.recovery_restart_count == 1
