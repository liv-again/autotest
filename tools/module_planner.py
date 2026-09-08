"""Build a module/page-group plan for row-scoped Excel UI execution.

The LLM reads a Sheet once and produces the planning context.  The executor
may reuse navigation inside one *navigation context*, but it still consumes
one Excel row at a time and records evidence for every row.  A page name alone
is deliberately not a grouping key: ``沪深A股-个股详情`` and
``港股-个股详情`` are different execution contexts.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Iterable

import yaml


CANONICAL_COLUMNS = {
    "一级目录": "level_1",
    "1级目录": "level_1",
    "二级目录": "level_2",
    "2级目录": "level_2",
    "三级目录": "level_3",
    "3级目录": "level_3",
    "四级目录": "level_4",
    "4级目录": "level_4",
    "用例名称": "case_name",
    "用例名": "case_name",
    "优先级": "priority",
    "入口": "entry",
    "步骤名称": "step_name",
    "步骤": "step_name",
    "前置条件": "precondition",
    "操作描述": "action",
    "操作": "action",
    "参数": "parameters",
    "预期结果": "expected",
    "预期": "expected",
}

LEVEL_FIELDS = ("level_1", "level_2", "level_3", "level_4")
GROUP_CONTEXT_FIELDS = LEVEL_FIELDS + ("entry",)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_key(value: Any) -> str:
    return "".join(_text(value).split()).lower()


def _column_map(headers: Iterable[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    aliases = {_header_key(key): value for key, value in CANONICAL_COLUMNS.items()}
    for index, header in enumerate(headers):
        canonical = aliases.get(_header_key(header))
        if canonical and canonical not in result:
            result[canonical] = index
    return result


def _normalise_group_part(value: Any) -> str:
    """Normalise a grouping component without losing its human-readable copy."""

    return " ".join(_text(value).split()).casefold()


def _carry_hierarchy(previous: dict[str, str], raw: dict[str, str]) -> dict[str, str]:
    """Expand merged Excel hierarchy cells and clear stale child levels.

    Excel commonly stores a merged ``一级目录`` cell only on its first row.
    When a new level is encountered, values from deeper levels belong to the
    previous branch and must not leak into the new navigation context.
    """

    context = dict(previous)
    for index, field in enumerate(LEVEL_FIELDS):
        value = raw.get(field, "")
        if not value:
            continue
        context[field] = value
        for child in LEVEL_FIELDS[index + 1 :]:
            context.pop(child, None)
    return {field: context.get(field, "") for field in LEVEL_FIELDS}


def _navigation_context(
    sheet_name: str,
    hierarchy: dict[str, str],
    *,
    entry: str,
    precondition: str,
    step_name: str,
) -> dict[str, Any]:
    values = {field: hierarchy.get(field, "") for field in LEVEL_FIELDS}
    values.update(
        {
            "sheet": sheet_name,
            "entry": entry,
            "precondition": precondition,
            "step_name": step_name,
        }
    )
    parts = [sheet_name, *(values[field] for field in LEVEL_FIELDS), entry]
    normalised = [_normalise_group_part(part) for part in parts]
    # Keep empty positions out of the key, but retain sheet as the first
    # component so identical level labels in two Sheets cannot be merged.
    key = "|".join(part for part in normalised if part)
    values["page_group_key"] = key or _normalise_group_part(sheet_name)
    values["display_path"] = " / ".join(
        part for part in (sheet_name, *(values[field] for field in LEVEL_FIELDS), entry) if part
    )
    return values


def _row_case(
    sheet_name: str,
    row: int,
    values: list[Any],
    columns: dict[str, int],
    source_order: int,
    hierarchy: dict[str, str] | None = None,
) -> tuple[dict[str, Any] | None, dict[str, str]]:
    def get(name: str) -> str:
        index = columns.get(name)
        return _text(values[index]) if index is not None and index < len(values) else ""

    raw_hierarchy = {field: get(field) for field in LEVEL_FIELDS}
    hierarchy = _carry_hierarchy(hierarchy or {}, raw_hierarchy)
    case_name = get("case_name")
    action = get("action")
    expected = get("expected")
    if not any((case_name, action, expected)):
        return None, hierarchy
    context = _navigation_context(
        sheet_name,
        hierarchy,
        entry=get("entry"),
        precondition=get("precondition"),
        step_name=get("step_name"),
    )
    case = {
        "case_id": f"{sheet_name}-row-{row:03d}",
        "module": sheet_name,
        "sheet": sheet_name,
        "row": row,
        "source_order": source_order,
        "case_name": case_name,
        "priority": get("priority"),
        "entry": get("entry"),
        "step_name": get("step_name"),
        "precondition": get("precondition"),
        "action": action,
        "parameters": get("parameters"),
        "expected": expected,
        **{field: hierarchy.get(field, "") for field in LEVEL_FIELDS},
        "navigation_context": context,
        "page_group_key": context["page_group_key"],
        "execution_unit": "single_excel_row",
    }
    return case, hierarchy


def _read_xls(source: Path, sheet_names: list[str]) -> list[dict[str, Any]]:
    import xlrd

    book = xlrd.open_workbook(str(source), formatting_info=False)
    available = book.sheet_names()
    selected = sheet_names or available
    missing = [name for name in selected if name not in available]
    if missing:
        raise ValueError(f"Excel 中不存在 Sheet: {', '.join(missing)}")
    cases: list[dict[str, Any]] = []
    source_order = 0
    for sheet_name in selected:
        sheet = book.sheet_by_name(sheet_name)
        columns = _column_map(sheet.row_values(0)) if sheet.nrows else {}
        hierarchy: dict[str, str] = {}
        for row_index in range(1, sheet.nrows):
            item, hierarchy = _row_case(
                sheet_name,
                row_index + 1,
                sheet.row_values(row_index),
                columns,
                source_order + 1,
                hierarchy,
            )
            if item:
                source_order += 1
                item["source_order"] = source_order
                cases.append(item)
    return cases


def _read_xlsx(source: Path, sheet_names: list[str]) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    book = load_workbook(source, read_only=True, data_only=True)
    available = list(book.sheetnames)
    selected = sheet_names or available
    missing = [name for name in selected if name not in available]
    if missing:
        raise ValueError(f"Excel 中不存在 Sheet: {', '.join(missing)}")
    cases: list[dict[str, Any]] = []
    source_order = 0
    for sheet_name in selected:
        sheet = book[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            continue
        columns = _column_map(headers)
        hierarchy: dict[str, str] = {}
        for row_number, values in enumerate(rows, start=2):
            item, hierarchy = _row_case(
                sheet_name,
                row_number,
                list(values),
                columns,
                source_order + 1,
                hierarchy,
            )
            if item:
                source_order += 1
                item["source_order"] = source_order
                cases.append(item)
    return cases


def load_cases(source: str | Path, sheet_names: Iterable[str] | None = None) -> list[dict[str, Any]]:
    path = Path(source).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(path)
    selected = list(sheet_names or [])
    if path.suffix.lower() == ".xls":
        return _read_xls(path, selected)
    return _read_xlsx(path, selected)


def _page_groups(module_name: str, cases: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Create contiguous page groups without reordering Excel rows.

    The same key appearing again after another key becomes a new group.  This
    preserves source order and prevents a later case from jumping backwards
    just because it happens to share a page label with an earlier case.
    """

    groups: list[dict[str, Any]] = []
    for case in cases:
        key = case.get("page_group_key") or module_name
        if not groups or groups[-1]["page_group_key"] != key:
            group_number = len(groups) + 1
            groups.append(
                {
                    "page_group_id": f"{module_name}-page-group-{group_number:03d}",
                    "page_group_order": group_number,
                    "page_group_key": key,
                    "navigation_context": dict(case.get("navigation_context") or {}),
                    "cases": [],
                }
            )
        group = groups[-1]
        case["page_group_id"] = group["page_group_id"]
        case["page_group_order"] = group["page_group_order"]
        group["cases"].append(case)
    for group in groups:
        group["expected_count"] = len(group["cases"])
        group["case_ids"] = [case["case_id"] for case in group["cases"]]
        group["row_range"] = [
            int(group["cases"][0]["row"]),
            int(group["cases"][-1]["row"]),
        ]
    return groups


def _profile_context(profile_path: str | Path | None) -> dict[str, Any]:
    if not profile_path:
        return {}
    path = Path(profile_path).expanduser().resolve()
    if not path.exists():
        return {"profile_file": str(path), "status": "missing"}
    try:
        document = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except (OSError, UnicodeDecodeError, yaml.YAMLError) as exc:
        return {"profile_file": str(path), "status": "unreadable", "error": str(exc)}
    if not isinstance(document, dict):
        return {"profile_file": str(path), "status": "invalid"}
    return {
        "profile_file": str(path),
        "status": "loaded",
        "slug": document.get("slug", ""),
        "app_version": document.get("app_version", ""),
        "entries": [
            {
                "key": item.get("key", ""),
                "path": item.get("path", ""),
                "status": item.get("status", ""),
                "app_version": item.get("app_version", ""),
            }
            for item in document.get("entries", [])
            if isinstance(item, dict)
        ],
        "capabilities": [
            {
                "key": item.get("key", ""),
                "supported": item.get("supported"),
                "note": item.get("note", ""),
                "status": item.get("status", ""),
            }
            for item in document.get("capabilities", [])
            if isinstance(item, dict)
        ],
    }


def _attach_profile_hints(groups: list[dict[str, Any]], profile: dict[str, Any]) -> None:
    entries = profile.get("entries") or []
    for group in groups:
        context = group.get("navigation_context") or {}
        source = " ".join(
            str(context.get(field, ""))
            for field in ("sheet", "level_1", "level_2", "level_3", "level_4", "entry", "display_path")
        ).casefold()
        tokens = [token for token in re.split(r"[^a-z0-9\u4e00-\u9fff]+", source) if len(token) >= 2]
        hints = []
        for entry in entries:
            path = _text(entry.get("path")).casefold()
            key = _text(entry.get("key")).casefold()
            if any(token in path or token in key for token in tokens):
                hints.append(entry)
        group["profile_hints"] = hints[:8]
        for case in group.get("cases", []):
            case["profile_hints"] = list(group["profile_hints"])


def build_module_plan(
    source: str | Path,
    sheet_names: Iterable[str] | None = None,
    *,
    profile_path: str | Path | None = None,
) -> dict[str, Any]:
    cases = load_cases(source, sheet_names)
    profile = _profile_context(profile_path)
    modules: list[dict[str, Any]] = []
    for module_name in dict.fromkeys(case["module"] for case in cases):
        module_cases = [case for case in cases if case["module"] == module_name]
        for module_order, case in enumerate(module_cases, start=1):
            case["module_order"] = module_order
        page_groups = _page_groups(module_name, module_cases)
        _attach_profile_hints(page_groups, profile)
        modules.append(
            {
                "module": module_name,
                "sheet": module_name,
                "expected_count": len(module_cases),
                "page_group_count": len(page_groups),
                "page_groups": page_groups,
                "cases": module_cases,
            }
        )
    return {
        "schema_version": "1.1",
        "plan_type": "module_page_group_execution_plan",
        "planning_scope": "module_page_group",
        "execution_scope": "single_excel_row",
        "review_scope": "module_exceptions",
        "page_batching_allowed": True,
        "grouping_strategy": "contiguous_navigation_context",
        "llm_instruction": (
            "LLM 读取模块用例和 App 画像后按导航上下文规划页面组；"
            "页面组只复用导航，业务动作、截图、断言和结果仍按 Excel 行独立执行。"
        ),
        "llm_roles": {
            "planner": "按一级至四级目录、入口和前置条件识别导航上下文，不能仅按目标页面名称合并用例。",
            "executor": "组内复用已验证入口；每行执行前后重新确认页面，异常行允许一次运行时重新规划。",
            "reviewer": "依据每行 action_trace、页面观察、预期结果和截图判断，不得用点击成功替代结果验证。",
            "profile_feedback": "只输出带证据的画像候选，未经 schema、版本和重复校验不得覆盖正式画像。",
        },
        "source_file": str(Path(source).expanduser().resolve()),
        "app_profile_context": {
            "profile_file": profile.get("profile_file", ""),
            "status": profile.get("status", "not_configured"),
            "slug": profile.get("slug", ""),
            "app_version": profile.get("app_version", ""),
            "entry_count": len(profile.get("entries") or []),
            "capability_count": len(profile.get("capabilities") or []),
        },
        "selected_sheets": [module["sheet"] for module in modules],
        "expected_count": len(cases),
        "execution_manifest": {
            "manifest_version": "2.0",
            "mode": "full",
            "execution_scope": "single_excel_row",
            "llm_review_required": True,
            "source_file": str(Path(source).expanduser().resolve()),
            "source_sheets": [module["sheet"] for module in modules],
            "expected_count": len(cases),
            "selected_cases": [
                {
                    "case_id": case["case_id"],
                    "sheet": case["sheet"],
                    "row": case["row"],
                    "source_order": case["source_order"],
                    "case_name": case["case_name"],
                    "page_group_id": case["page_group_id"],
                    "page_group_key": case["page_group_key"],
                    "navigation_context": case["navigation_context"],
                }
                for case in cases
            ],
            "page_groups": [
                {
                    "page_group_id": group["page_group_id"],
                    "page_group_key": group["page_group_key"],
                    "case_ids": group["case_ids"],
                    "row_range": group["row_range"],
                }
                for module in modules
                for group in module["page_groups"]
            ],
        },
        "modules": modules,
    }


def write_plan(plan: dict[str, Any], output: str | Path) -> Path:
    path = Path(output).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(plan, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def main() -> int:
    parser = argparse.ArgumentParser(description="按 Sheet 生成模块规划，供逐行执行器消费")
    parser.add_argument("--source", required=True, help=".xls/.xlsx 用例文件")
    parser.add_argument("--sheet", action="append", dest="sheets", help="指定 Sheet，可重复；默认读取全部")
    parser.add_argument("--profile", help="App profile.yaml，为页面组规划提供入口和版本上下文")
    parser.add_argument("--out", required=True, help="module_plan.json 输出路径")
    args = parser.parse_args()
    plan = build_module_plan(args.source, args.sheets, profile_path=args.profile)
    output = write_plan(plan, args.out)
    print(json.dumps({"out": str(output), "modules": len(plan["modules"]), "cases": plan["expected_count"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
